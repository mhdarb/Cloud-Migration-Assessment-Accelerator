"""Client questionnaires: upload -> answer from evidence -> download with answers written in.

Covers every accepted format's parse + write-back round trip, spreadsheet formula-injection
safety, the answer cache, the busy/not-run gates, isolation from the Questions tab, and
the HTTP flow end to end.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi import UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import get_db
from app.models.entities import (
    Claim,
    PipelineStatus,
    Questionnaire,
    QuestionnaireItem,
    ReviewDecision,
)
from app.services import questionnaires as svc
from app.services.assessment_questions import build_assessment_answers
from app.services.questionnaire_io import (
    AnswerCell,
    clean_question,
    looks_like_question,
    parse_questionnaire,
    write_answered,
)

EVIL = '=HYPERLINK("http://evil.example","click")'


def _cell(answer: str = "Yes, nightly.", *, review: bool = False) -> AnswerCell:
    return AnswerCell(answer=answer, confidence=0.82, needs_review=review, sources="runbook.docx (Q3)")


def _answer_all(path: Path, fmt: str, answer: str = "Yes, nightly.", *, summary: bool = False) -> tuple[bytes, str]:
    items = parse_questionnaire(str(path), fmt, path.name)
    return write_answered(fmt, str(path), [(item, _cell(answer)) for item in items], summary=summary)


# --------------------------------------------------------------------------- #
# Question detection
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("raw", "clean"),
    [
        ("Q1: What is the RPO?", "What is the RPO?"),
        ("1.2) Describe the backup process", "Describe the backup process"),
        ("Question 3. Which databases hold PII?", "Which databases hold PII?"),
        ("a) Is TLS enforced?", "Is TLS enforced?"),
        ("A3. Describe how secrets are managed", "Describe how secrets are managed"),
        ("SEC-12: Is MFA enforced for admins?", "Is MFA enforced for admins?"),
        ("What is the RPO?", "What is the RPO?"),
    ],
)
def test_clean_question_strips_numbering(raw, clean):
    assert clean_question(raw) == clean


def test_question_detection():
    assert looks_like_question("Is data encrypted at rest?")
    assert looks_like_question("4. Describe the disaster recovery process")  # numbered imperative
    assert not looks_like_question("Describe the process")  # unnumbered prose, no '?'
    assert not looks_like_question("Section 2 - Security")
    assert not looks_like_question("Why?")  # too short to be a real question
    assert looks_like_question("Backup retention period and location", in_question_column=True)
    assert not looks_like_question("Security", in_question_column=True)  # section label


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _workbook(path: Path, rows: list[list[object]], title: str = "Questionnaire") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_xlsx_fills_existing_answer_column(tmp_path):
    path = _workbook(
        tmp_path / "rfi.xlsx",
        [
            ["Client RFI — Zephyr"],
            [],
            ["#", "Question", "Answer", "Owner"],
            [1, "What is the RPO for the order database?", None, "DBA"],
            [2, "Security"],  # a section label in the question column: not a question
            [3, "Is TLS 1.2 enforced on all public endpoints?", None, "SecOps"],
        ],
    )
    items = parse_questionnaire(str(path), "xlsx")
    assert [i.question for i in items] == [
        "What is the RPO for the order database?",
        "Is TLS 1.2 enforced on all public endpoints?",
    ]
    assert items[0].locator == {
        "kind": "xlsx",
        "sheet": "Questionnaire",
        "row": 4,
        "col": 2,
        "answer_col": 3,
        "header_row": 3,
    }

    data, suffix = _answer_all(path, "xlsx")
    assert suffix == ".xlsx"
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["C4"].value == "Yes, nightly." and ws["C6"].value == "Yes, nightly."
    assert ws["D4"].value == "DBA"  # the client's own columns are untouched
    assert [ws.cell(row=3, column=c).value for c in (5, 6, 7)] == ["Confidence", "Needs review", "Sources"]
    assert ws["E4"].value == "82%" and ws["F4"].value == "No" and ws["G4"].value == "runbook.docx (Q3)"
    assert ws["C5"].value is None  # the section-label row gets no answer


def test_xlsx_without_answer_column_appends_one(tmp_path):
    path = _workbook(tmp_path / "q.xlsx", [["Question"], ["Which apps are in scope?"], ["Who owns DR?"]])
    ws = load_workbook(io.BytesIO(_answer_all(path, "xlsx")[0])).active
    assert [c.value for c in ws[1]] == ["Question", "Answer", "Confidence", "Needs review", "Sources"]
    assert ws["B2"].value == "Yes, nightly."


def test_xlsx_headerless_question_column_is_detected(tmp_path):
    path = _workbook(
        tmp_path / "loose.xlsx",
        [["S1", "What is the peak TPS?"], ["S2", "Is there a DR site?"], ["S3", "Which OS versions run?"]],
    )
    items = parse_questionnaire(str(path), "xlsx")
    assert len(items) == 3 and all(i.locator["col"] == 2 for i in items)


def test_xlsx_answer_can_never_become_a_formula(tmp_path):
    path = _workbook(tmp_path / "q.xlsx", [["Question", "Answer"], ["Where is the admin portal?", None]])
    ws = load_workbook(io.BytesIO(_answer_all(path, "xlsx", EVIL)[0])).active
    assert ws["B2"].value == EVIL
    assert ws["B2"].data_type == "s"  # stored as text, not "f"


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def test_csv_round_trip_keeps_delimiter_and_escapes_formulas(tmp_path):
    path = tmp_path / "q.csv"
    path.write_text("ID;Question;Answer\n1;What is the RTO?;\n2;Is MFA enforced?;\n", encoding="utf-8")
    items = parse_questionnaire(str(path), "csv")
    assert [i.question for i in items] == ["What is the RTO?", "Is MFA enforced?"]

    data, suffix = write_answered("csv", str(path), [(items[0], _cell("4 hours")), (items[1], _cell(EVIL))])
    assert suffix == ".csv"
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["ID", "Question", "Answer", "Confidence", "Needs review", "Sources"]
    assert rows[1][:3] == ["1", "What is the RTO?", "4 hours"]
    assert rows[2][2] == "'" + EVIL  # neutralised for Excel


def test_csv_without_answer_column(tmp_path):
    path = tmp_path / "q.csv"
    path.write_text("Question\nWhich regions are approved?\n", encoding="utf-8")
    rows = list(csv.reader(io.StringIO(_answer_all(path, "csv")[0].decode("utf-8-sig"))))
    assert rows[0] == ["Question", "Answer", "Confidence", "Needs review", "Sources"]
    assert rows[1][1] == "Yes, nightly."


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def test_docx_paragraph_questions_answered_in_place(tmp_path):
    doc = DocxDocument()
    doc.add_heading("Migration questionnaire", 1)
    doc.add_paragraph("Please answer every question below.")
    doc.add_paragraph("1. What is the backup schedule?")
    doc.add_paragraph("Answer:")  # the template's own blank slot is filled, not duplicated
    doc.add_paragraph("2. Describe the DR failover process")
    doc.add_paragraph("Closing notes.")
    path = tmp_path / "q.docx"
    doc.save(path)

    items = parse_questionnaire(str(path), "docx")
    assert [i.question for i in items] == ["What is the backup schedule?", "Describe the DR failover process"]
    assert items[0].locator["answer_paragraph"] == 3

    data, suffix = _answer_all(path, "docx")
    assert suffix == ".docx"
    texts = [p.text for p in DocxDocument(io.BytesIO(data)).paragraphs]
    q1 = texts.index("1. What is the backup schedule?")
    assert texts[q1 + 1] == "Answer: Yes, nightly."
    assert texts[q1 + 2].startswith("Confidence 82%")
    q2 = texts.index("2. Describe the DR failover process")
    assert texts[q2 + 1] == "Answer: Yes, nightly."
    assert texts[-1] == "Closing notes."
    assert "Answer:" not in texts  # no leftover blank slot


def test_docx_table_questionnaire(tmp_path):
    doc = DocxDocument()
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [["#", "Question", "Response"], ["1", "Is data encrypted at rest?", ""], ["2", "Who approves changes?", ""]]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    path = tmp_path / "t.docx"
    doc.save(path)

    items = parse_questionnaire(str(path), "docx")
    assert [i.locator["kind"] for i in items] == ["docx_table", "docx_table"]
    out = DocxDocument(io.BytesIO(_answer_all(path, "docx")[0])).tables[0]
    assert out.cell(1, 2).text.startswith("Answer: Yes, nightly.")
    assert out.cell(1, 1).text == "Is data encrypted at rest?"


# --------------------------------------------------------------------------- #
# Text, Markdown, PDF
# --------------------------------------------------------------------------- #
def test_markdown_answers_inserted_after_each_question(tmp_path):
    path = tmp_path / "q.md"
    path.write_text("# RFI\n\n## Q1. What is the RPO?\n\nSome context.\n\n- Is MFA enforced?\n", encoding="utf-8")
    text = _answer_all(path, "md")[0].decode()
    assert "## Q1. What is the RPO?\n\n**Answer:** Yes, nightly." in text
    assert "- Is MFA enforced?\n\n**Answer:** Yes, nightly." in text
    assert "Some context." in text


def test_txt_fills_answer_slot(tmp_path):
    path = tmp_path / "q.txt"
    path.write_text("Q1: What is the RTO?\nAnswer:\nQ2: Who is the app owner?\n", encoding="utf-8")
    lines = _answer_all(path, "txt")[0].decode().splitlines()
    assert lines[:2] == ["Q1: What is the RTO?", "Answer: Yes, nightly."]
    assert "Answer:" not in lines


def test_pdf_questionnaire_exports_excel_summary(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "rfi.pdf"
    c = canvas.Canvas(str(path))
    for y, line in (
        (800, "Vendor security questionnaire"),
        (770, "1. Is data encrypted at rest?"),
        (750, "2. What is the RPO?"),
    ):
        c.drawString(72, y, line)
    c.save()

    items = parse_questionnaire(str(path), "pdf", "rfi.pdf")
    assert [i.question for i in items] == ["Is data encrypted at rest?", "What is the RPO?"]
    data, suffix = _answer_all(path, "pdf")
    assert suffix == ".xlsx"
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["#", "Question", "Answer", "Confidence", "Needs review", "Sources"]
    assert ws["B2"].value == "Is data encrypted at rest?" and ws["C2"].value == "Yes, nightly."


def test_duplicate_questions_are_asked_once(tmp_path):
    path = tmp_path / "q.txt"
    path.write_text("What is the RPO?\nIs MFA on?\nQ7. What is the RPO ?\n", encoding="utf-8")
    assert [i.question for i in parse_questionnaire(str(path), "txt")] == ["What is the RPO?", "Is MFA on?"]


# --------------------------------------------------------------------------- #
# Service
# --------------------------------------------------------------------------- #
def _upload(path: Path) -> UploadFile:
    return UploadFile(file=io.BytesIO(path.read_bytes()), filename=path.name)


def _completed(db, assessment) -> None:
    from datetime import datetime

    assessment.status = PipelineStatus.completed
    assessment.pipeline_finished_at = datetime(2026, 9, 1, 12, 0, 0)
    db.add(
        Claim(
            assessment_id=assessment.id,
            entity_type="database",
            entity_key="orders-db",
            attribute="rpo",
            value="15 minutes",
            confidence=0.9,
            evidence_refs=["c-rpo"],
            evidence_quote="orders-db RPO is 15 minutes",
            unsupported=False,
            is_selected=True,
            needs_human_review=False,
        )
    )
    db.commit()


class _CountingProse:
    def __init__(self) -> None:
        self.calls = 0

    def rewrite_question_answers(self, answers) -> None:
        self.calls += 1


async def _import(db, assessment, tmp_path, text="Q1: What is the RPO?\nQ2: Which colour is the logo xyzzy?\n"):
    path = tmp_path / "client.txt"
    path.write_text(text, encoding="utf-8")
    return await svc.import_questionnaire(db, assessment, _upload(path))


@pytest.mark.anyio
async def test_import_answer_and_cache(db_session, assessment, tmp_path):
    q = await _import(db_session, assessment, tmp_path)
    assert svc.questionnaire_summary(db_session, q)["question_count"] == 2

    with pytest.raises(ValueError, match="Run the assessment"):
        svc.questionnaire_answers(db_session, assessment, q)
    assessment.status = PipelineStatus.extracting
    db_session.commit()
    with pytest.raises(svc.QuestionnaireBusy):
        svc.questionnaire_answers(db_session, assessment, q)

    _completed(db_session, assessment)
    prose = _CountingProse()
    answers = svc.questionnaire_answers(db_session, assessment, q, prose=prose)
    assert [a["position"] for a in answers] == [1, 2]
    assert "15 minutes" in answers[0]["answer"] and answers[0]["supported"] is True
    assert answers[1]["supported"] is False and answers[1]["needs_human_review"] is True

    # Preview then download: one computation.
    svc.export_questionnaire(db_session, assessment, q, prose=prose)
    assert prose.calls == 1
    # A review decision can change answers, so it invalidates the cache.
    db_session.add(ReviewDecision(assessment_id=assessment.id, target_type="claim", target_key="k", action="accept"))
    db_session.commit()
    svc.questionnaire_answers(db_session, assessment, q, prose=prose)
    assert prose.calls == 2


@pytest.mark.anyio
async def test_questionnaire_questions_stay_out_of_the_questions_tab(db_session, assessment, tmp_path):
    await _import(db_session, assessment, tmp_path)
    _completed(db_session, assessment)
    origins = {a["origin"] for a in build_assessment_answers(db_session, assessment.id)["answers"]}
    assert origins == {"standard"}


@pytest.mark.anyio
async def test_rejected_uploads(db_session, assessment, tmp_path):
    legacy = tmp_path / "old.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match="Save it as .xlsx"):
        await svc.import_questionnaire(db_session, assessment, _upload(legacy))
    with pytest.raises(ValueError, match="No questions found"):
        await _import(db_session, assessment, tmp_path, text="Just some notes.\nNothing to answer here.\n")
    assert db_session.query(Questionnaire).count() == 0
    stored = Path(tmp_path / "uploads" / assessment.id)
    assert not stored.exists() or not any(stored.iterdir())  # the rejected file isn't kept


@pytest.mark.anyio
async def test_delete_removes_rows_and_file(db_session, assessment, tmp_path):
    q = await _import(db_session, assessment, tmp_path)
    stored = Path(q.storage_path)
    assert stored.exists()
    svc.delete_questionnaire(db_session, q)
    assert db_session.query(QuestionnaireItem).count() == 0
    assert not stored.exists()


@pytest.fixture
def anyio_backend():
    return "asyncio"


# --------------------------------------------------------------------------- #
# HTTP
# --------------------------------------------------------------------------- #
@pytest.fixture
def shared_session(db_session):
    """TestClient runs sync endpoints on a worker thread, and in-memory SQLite gives every
    thread its own empty database — so the HTTP test needs one connection shared by all."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from app.database import Base

    engine = create_engine(
        "sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine, autoflush=False)()
    yield session
    session.close()
    engine.dispose()


def test_http_flow(shared_session, tmp_path):
    from app.main import app
    from app.models.entities import Assessment

    db_session = shared_session
    assessment = Assessment(name="HTTP", status=PipelineStatus.pending)
    db_session.add(assessment)
    db_session.commit()
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        client = TestClient(app)
        base = f"/assessments/{assessment.id}/questionnaires"
        xlsx = _workbook(
            tmp_path / "Client RFI.xlsx", [["Question", "Answer"], ["What is the RPO for orders-db?", None]]
        )
        created = client.post(base, files={"file": (xlsx.name, xlsx.read_bytes())})
        assert created.status_code == 200, created.text
        body = created.json()
        assert body["question_count"] == 1 and body["answered_format"] == "xlsx"
        assert [q["id"] for q in client.get(base).json()] == [body["id"]]

        assert client.get(f"{base}/{body['id']}/answers").status_code == 400  # not run yet
        _completed(db_session, assessment)
        preview = client.get(f"{base}/{body['id']}/answers").json()
        assert "15 minutes" in preview["answers"][0]["answer"]

        download = client.get(f"{base}/{body['id']}/download")
        assert download.status_code == 200
        assert 'filename="Client RFI - answered.xlsx"' in download.headers["content-disposition"]
        assert load_workbook(io.BytesIO(download.content)).active["B2"].value == preview["answers"][0]["answer"]
        assert client.get(f"{base}/{body['id']}/download?format=bogus").status_code == 400

        bad = client.post(base, files={"file": ("notes.txt", b"nothing to see")})
        assert bad.status_code == 400 and "No questions found" in bad.json()["detail"]

        assert client.delete(f"{base}/{body['id']}").status_code == 204
        assert client.get(f"{base}/{body['id']}/answers").status_code == 404
    finally:
        app.dependency_overrides.clear()
