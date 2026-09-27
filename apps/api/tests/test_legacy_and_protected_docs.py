"""Legacy Office conversion, irregular spreadsheet headers, and password-protected PDFs."""

from __future__ import annotations

import io
import os
import stat
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.config import get_settings
from app.models.entities import Document
from app.services import converters
from app.services.assessment_service import WrongPassword, unlock_document
from app.services.ingest import ingest_documents
from app.services.parsers import EncryptedDocument, parse_file

OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


# --------------------------------------------------------------------------- #
# Irregular / stacked spreadsheet headers
# --------------------------------------------------------------------------- #
def _xlsx(tmp_path: Path, rows: list[list], *, merges: tuple[str, ...] = (), name: str = "t.xlsx", **sheets) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Servers"
    for row in rows:
        ws.append(row)
    for rng in merges:
        ws.merge_cells(rng)
    for title, (state, extra) in sheets.items():
        other = wb.create_sheet(title)
        other.sheet_state = state
        for row in extra:
            other.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


def _lines(path: Path) -> list[str]:
    return parse_file(str(path), path.name).pages[0].text.splitlines()


def test_group_band_over_labels_becomes_one_header(tmp_path):
    path = _xlsx(
        tmp_path,
        [
            ["Estate export"],
            [],
            ["Identity", None, "Capacity", None, "Utilisation"],
            ["hostname", "os", "vCPU", "RAM (GB)", "CPU %"],
            ["app-01", "RHEL 9", 4, 16, 55],
            ["app-02", "RHEL 9", 8, 32, 61],
        ],
        merges=("A3:B3", "C3:D3"),
    )
    lines = _lines(path)
    assert lines[0] == "# Sheet: Servers | Estate export"
    assert lines[1] == "hostname | os | vCPU | RAM (GB) | CPU %"
    assert lines[2].startswith("app-01 | RHEL 9 | 4")
    assert not any(line.startswith("Identity") for line in lines)


def test_duplicate_leaf_labels_take_their_group_as_prefix(tmp_path):
    path = _xlsx(
        tmp_path,
        [
            ["host", "CPU", None, "Memory", None],
            [None, "Avg", "Peak", "Avg", "Peak"],
            ["app-01", 30, 70, 40, 80],
        ],
        merges=("B1:C1", "D1:E1"),
    )
    assert _lines(path)[1] == "host | CPU Avg | CPU Peak | Memory Avg | Memory Peak"


def test_restated_header_row_is_merged_not_treated_as_data(tmp_path):
    path = _xlsx(
        tmp_path, [["device", "mgmt_ip", "vlan"], ["Device", "Management IP", "VLAN"], ["lb-01", "10.0.0.1", 20]]
    )
    assert _lines(path)[1:] == ["device | mgmt_ip | vlan", "lb-01 | 10.0.0.1 | 20"]


def test_repeated_header_inside_data_totals_and_trailing_notes(tmp_path):
    path = _xlsx(
        tmp_path,
        [
            ["hostname", "vcpu", "memory_gb"],
            ["app-01", 4, 16],
            ["hostname", "vcpu", "memory_gb"],  # page-break repeat in a printed export
            ["app-02", 8, 32],
            ["Subtotal", 12, 48],
            ["Grand Total", 12, 48],
            ["Note: figures from the Q3 capacity review, not live data."],
        ],
    )
    lines = _lines(path)
    assert lines[1:4] == ["hostname | vcpu | memory_gb", "app-01 | 4 | 16", "app-02 | 8 | 32"]
    assert "# NOTE: Note: figures from the Q3 capacity review, not live data." in lines
    assert "# NOTE: 2 total/subtotal row(s) excluded from the table" in lines
    assert "# NOTE: 1 repeated header row(s) removed" in lines


def test_all_text_table_keeps_its_first_data_row(tmp_path):
    """A label-only first data row must not be swallowed as a second header when the
    following rows are text too (e.g. a lookup or owner list)."""
    path = _xlsx(tmp_path, [["team", "owner"], ["Payments", "Tom"], ["Fleet", "Ana"]])
    assert _lines(path)[1:] == ["team | owner", "Payments | Tom", "Fleet | Ana"]


def test_hidden_sheets_are_skipped_with_a_warning_and_can_be_included(tmp_path, monkeypatch):
    path = _xlsx(tmp_path, [["hostname", "vcpu"], ["app-01", 4]], Lookups=("hidden", [["status"], ["Retired"]]))
    result = parse_file(str(path), path.name)
    assert [p.label for p in result.pages] == ["Servers"]
    assert any("Hidden sheet(s) 'Lookups' skipped" in w for w in result.warnings)
    monkeypatch.setenv("XLSX_INCLUDE_HIDDEN_SHEETS", "true")
    get_settings.cache_clear()
    assert [p.label for p in parse_file(str(path), path.name).pages] == ["Servers", "Lookups"]


def test_formula_cells_without_saved_results_are_reported(tmp_path):
    path = _xlsx(tmp_path, [["hostname", "vcpu"], ["app-01", 4], ["app-02", 8], ["TOTAL", "=SUM(B2:B3)"]])
    result = parse_file(str(path), path.name)
    assert any("1 formula cell(s) with no saved result" in w for w in result.warnings)
    assert not any(line.startswith("TOTAL") for line in result.pages[0].text.splitlines())


def test_caption_line_is_never_taken_as_the_inventory_header():
    """The ShopFront bug: '# Sheet: ... | ...' was read as the header, dropping every host."""
    from app.models.entities import Chunk
    from app.services.heuristic_extract import _extract_inventory_columns

    chunk = Chunk(
        id="c1",
        text="# Sheet: Infrastructure Inventory | ShopFront — Hosts | Production\n"
        "Hostname | Role | vCPU | RAM (GB)\nshopfront-web-p01 | Web | 4 | 16\n",
    )
    claims = _extract_inventory_columns(chunk)
    assert {(c.entity_key, c.attribute, c.value) for c in claims} >= {
        ("shopfront-web-p01", "vcpus", "4"),
        ("shopfront-web-p01", "memory_gb", "16"),
    }


# --------------------------------------------------------------------------- #
# Legacy .doc / .xls
# --------------------------------------------------------------------------- #
def _legacy(tmp_path: Path, name: str, extra: bytes = b"") -> Path:
    path = tmp_path / name
    path.write_bytes(OLE2 + extra + b"\x00" * 512)
    return path


def test_legacy_doc_without_libreoffice_explains_how_to_enable_it(tmp_path, monkeypatch):
    monkeypatch.setattr(converters, "libreoffice_binary", lambda: None)
    with pytest.raises(ValueError, match="legacy binary .doc format needs LibreOffice"):
        parse_file(str(_legacy(tmp_path, "spec.doc")), "spec.doc")


def test_legacy_doc_is_converted_then_parsed_as_docx(tmp_path, monkeypatch):
    modern = tmp_path / "modern.docx"
    doc = DocxDocument()
    doc.add_paragraph("Legacy spec: Billing API runs on app-01.")
    doc.save(modern)
    seen = {}

    def fake_convert(path, target_ext, *, kind):
        seen.update(path=path, target=target_ext, kind=kind)
        return str(modern)

    monkeypatch.setattr(converters, "convert_legacy_office", fake_convert)
    result = parse_file(str(_legacy(tmp_path, "spec.doc")), "spec.doc")
    assert seen["target"] == "docx" and seen["kind"] == ".doc"
    assert "Billing API runs on app-01" in result.pages[0].text
    assert any("Converted legacy .doc" in w for w in result.warnings)


def test_legacy_xls_is_converted_then_parsed_as_xlsx(tmp_path, monkeypatch):
    modern = _xlsx(tmp_path, [["hostname", "vcpu"], ["app-01", 4]], name="modern.xlsx")
    monkeypatch.setattr(converters, "convert_legacy_office", lambda path, target_ext, *, kind: str(modern))
    result = parse_file(str(_legacy(tmp_path, "cmdb.xls")), "cmdb.xls")
    assert "app-01 | 4" in result.pages[0].text


def test_password_protected_office_file_is_reported_not_converted(tmp_path, monkeypatch):
    monkeypatch.setattr(converters, "convert_legacy_office", lambda *a, **k: pytest.fail("must not convert"))
    path = _legacy(tmp_path, "locked.xlsx", "EncryptedPackage".encode("utf-16-le"))
    with pytest.raises(EncryptedDocument, match="password-protected Office document"):
        parse_file(str(path), "locked.xlsx")


def test_libreoffice_command_and_cache(tmp_path, monkeypatch):
    """Drive convert_legacy_office with a stand-in `soffice` that writes its output the
    way LibreOffice does, so the real command line and caching logic are exercised."""
    fake = tmp_path / "soffice"
    log = tmp_path / "calls.log"
    fake.write_text(
        "#!/bin/sh\n"
        f'echo "$@" >> "{log}"\n'
        'while [ "$1" != "--outdir" ]; do shift; done\n'
        'outdir="$2"; src="$3"; stem=$(basename "$src"); stem="${stem%.*}"\n'
        'printf converted > "$outdir/$stem.docx"\n'
    )
    fake.chmod(fake.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("LIBREOFFICE_PATH", str(fake))
    get_settings.cache_clear()
    source = _legacy(tmp_path, "spec.doc")

    first = converters.convert_legacy_office(str(source), "docx", kind=".doc")
    again = converters.convert_legacy_office(str(source), "docx", kind=".doc")
    assert first == again == str(source) + ".converted.docx"
    assert Path(first).read_text() == "converted"
    calls = log.read_text().splitlines()
    assert len(calls) == 1  # second call served from the cache beside the original
    assert "--headless" in calls[0] and "--convert-to docx" in calls[0]
    assert "-env:UserInstallation=file://" in calls[0]  # private, throwaway profile


def test_failed_conversion_is_a_clear_error(tmp_path, monkeypatch):
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\nexit 1\n")
    fake.chmod(fake.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("LIBREOFFICE_PATH", str(fake))
    get_settings.cache_clear()
    with pytest.raises(ValueError, match="could not convert legacy binary .doc"):
        converters.convert_legacy_office(str(_legacy(tmp_path, "bad.doc")), "docx", kind=".doc")


# --------------------------------------------------------------------------- #
# Password-protected PDFs
# --------------------------------------------------------------------------- #
def _pdf(path: Path, text: str, *, user: str | None = None, owner: str | None = None) -> Path:
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, text)
    c.showPage()
    c.save()
    buf.seek(0)
    writer = PdfWriter(clone_from=PdfReader(buf))
    if user is not None or owner is not None:
        writer.encrypt(user_password=user or "", owner_password=owner or "owner", permissions_flag=0)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def test_permission_restricted_pdf_is_read(tmp_path):
    path = _pdf(tmp_path / "sow.pdf", "Exit clause: 90 days notice", owner="vendor-owner")
    result = parse_file(str(path), "sow.pdf")
    assert "Exit clause" in result.pages[0].text
    assert any("permission restrictions" in w for w in result.warnings)


def test_open_password_pdf_needs_its_password(tmp_path):
    path = _pdf(tmp_path / "plan.pdf", "Capacity plan", user="s3cret")
    with pytest.raises(EncryptedDocument, match="use Unlock"):
        parse_file(str(path), "plan.pdf")


def _attach(db, assessment, path: Path) -> Document:
    doc = Document(
        assessment_id=assessment.id, filename=path.name, content_type="application/pdf", storage_path=str(path)
    )
    db.add(doc)
    db.commit()
    return doc


def test_ingest_flags_the_document_for_unlocking(db_session, assessment, tmp_path):
    doc = _attach(db_session, assessment, _pdf(tmp_path / "plan.pdf", "Capacity plan", user="s3cret"))
    result = ingest_documents(db_session, assessment.id)
    db_session.refresh(doc)
    assert doc.parse_summary["needs_password"] is True
    assert any("password-protected" in g for e in result.manifest_extractions for g in e.gaps)


def test_unlock_with_the_right_password_makes_it_readable(db_session, assessment, tmp_path):
    original = _pdf(tmp_path / "plan.pdf", "Capacity plan: pay-db-01 needs 64 GB", user="s3cret")
    doc = _attach(db_session, assessment, original)
    ingest_documents(db_session, assessment.id)

    with pytest.raises(WrongPassword):
        unlock_document(db_session, assessment, doc.id, "guess")
    unlock_document(db_session, assessment, doc.id, "s3cret")

    db_session.refresh(doc)
    assert not original.exists() and doc.storage_path.endswith(".unlocked.pdf")
    assert "needs_password" not in (doc.parse_summary or {})
    assert "pay-db-01 needs 64 GB" in parse_file(doc.storage_path, doc.filename).pages[0].text
    assert "s3cret" not in repr(doc.parse_summary) and "s3cret" not in doc.storage_path


def test_unlocking_an_unprotected_pdf_is_refused(db_session, assessment, tmp_path):
    doc = _attach(db_session, assessment, _pdf(tmp_path / "open.pdf", "hello"))
    with pytest.raises(ValueError, match="isn't password-protected"):
        unlock_document(db_session, assessment, doc.id, "anything")


@pytest.fixture(autouse=True)
def _reset_settings():
    yield
    for var in ("LIBREOFFICE_PATH", "XLSX_INCLUDE_HIDDEN_SHEETS"):
        os.environ.pop(var, None)
    get_settings.cache_clear()
