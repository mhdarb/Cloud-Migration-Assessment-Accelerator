"""Regression over the generated 'stress' estate (`sample-data/stress`).

These fixtures deliberately cover awkward real-world shapes (scanned PDF, merged data
cells, ragged CSV, nested-key JSON, markdown/embedded tables, legacy .doc, multilingual
text). The invariant this suite protects is the one that makes ingest trustworthy:

    every document either yields usable structured chunks, or degrades to a VISIBLE
    signal (a parse rejection or a surfaced warning) — never silent data loss.

The corpus is regenerated into a tmp dir from `scripts/generate_stress_fixtures.py`, so
the test is hermetic and never depends on committed binary fixtures.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

from app.services.chunkers import DocumentChunker
from app.services.parsers import TABLE_BLOCK_START, parse_file

_GEN = Path(__file__).resolve().parents[3] / "scripts" / "generate_stress_fixtures.py"


@pytest.fixture(scope="module")
def estate(tmp_path_factory) -> Path:
    """Generate the stress corpus into a temp dir and return it."""
    if not _GEN.exists():  # pragma: no cover - defensive
        pytest.skip("generator script not found")
    spec = importlib.util.spec_from_file_location("_stress_gen", _GEN)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # @dataclass resolves annotations via sys.modules
    spec.loader.exec_module(module)
    out = tmp_path_factory.mktemp("stress")
    module.OUT = out
    try:
        module.main()
    except ModuleNotFoundError as exc:  # pragma: no cover - reportlab/PIL missing
        pytest.skip(f"fixture generator dependency missing: {exc}")
    return out


def _parse(estate: Path, name: str):
    return parse_file(str(estate / name), name)


def test_no_document_is_silently_empty(estate: Path):
    """Every generated file either parses to real chunks, or is rejected, or carries a
    warning. Nothing is accepted as zero content with no trace."""
    for path in sorted(estate.iterdir()):
        if path.name == "README.md" or not path.is_file():
            continue
        try:
            result = parse_file(str(path), path.name)
        except Exception:
            continue  # a parse error is a visible signal — ingest surfaces it as a gap
        pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename=path.name)
        chars = sum(len(p.text.strip()) for p in result.pages)
        assert pieces or result.warnings or chars == 0 and result.warnings, (
            f"{path.name} produced neither chunks nor a warning"
        )
        if chars < 8:
            assert result.warnings, f"{path.name} is near-empty but surfaced no warning"


def test_scanned_pdf_warns(estate: Path):
    result = _parse(estate, "scanned-runbook.pdf")
    assert result.warnings
    assert "scanned" in result.warnings[0].lower() or "no extractable text" in result.warnings[0].lower()


def test_empty_text_warns(estate: Path):
    assert _parse(estate, "empty-notes.txt").warnings


def test_protected_pdf_is_rejected(estate: Path):
    """A password-protected PDF can't be extracted — it must raise (which ingest turns
    into an unreadable-document gap), never parse to silent empty content."""
    with pytest.raises(Exception):
        _parse(estate, "protected-capacity-plan.pdf")


def test_legacy_doc_rejected(estate: Path):
    with pytest.raises(ValueError, match="legacy binary .doc"):
        _parse(estate, "legacy-spec.doc")


def test_ruled_pdf_table_recovered(estate: Path):
    joined = "\n".join(p.text for p in _parse(estate, "architecture-brief.pdf").pages)
    assert TABLE_BLOCK_START in joined
    assert "ord-api-01 | Order API | RHEL 9 | 8 | 32" in joined


def test_two_column_pdf_not_mangled_into_table(estate: Path):
    """The two-column report must read as prose, not be carved into a bogus table."""
    result = _parse(estate, "capacity-report.pdf")
    joined = "\n".join(p.text for p in result.pages)
    assert TABLE_BLOCK_START not in joined
    assert "Order Management" in joined and "Fleet Tracking" in joined


def test_xlsx_merged_data_cells_and_summary(estate: Path):
    result = _parse(estate, "cmdb-export.xlsx")
    servers = result.pages[0].text
    # Merged 'datacenter' value propagated to every row, not just the anchor.
    assert servers.count("DC-North") >= 3
    assert servers.count("DC-South") >= 2
    pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename="cmdb-export.xlsx")
    assert any(p.metadata.get("kind") == "table_summary" for p in pieces)
    # The Network sheet's hidden second header trips the shape guard -> visible fallback.
    assert any(p.metadata.get("shape_guard_failed") for p in pieces)


def test_ragged_csv_hits_shape_guard(estate: Path):
    result = _parse(estate, "messy-inventory.csv")
    pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename="messy-inventory.csv")
    assert pieces  # not lost
    assert any(p.metadata.get("shape_guard_failed") for p in pieces)


def test_json_records_under_nested_key(estate: Path):
    text = _parse(estate, "cloud-assets.json").pages[0].text
    assert "asset_id | name | vcpu | memory_gb | os" in text
    assert "vm-101 | ord-api-01 | 8 | 32 | RHEL 9" in text


def test_markdown_table_becomes_block(estate: Path):
    text = _parse(estate, "migration-runbook.md").pages[0].text
    assert TABLE_BLOCK_START in text
    assert "ord-api-01 | primary | DC-North | 1" in text


def test_docx_embedded_table_chunked(estate: Path):
    result = _parse(estate, "solution-architecture.docx")
    pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename="solution-architecture.docx")
    assert any(p.metadata.get("embedded_table") for p in pieces)


def test_questionnaire_pairs(estate: Path):
    result = _parse(estate, "discovery-questionnaire.docx")
    pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename="discovery-questionnaire.docx")
    assert sum(1 for p in pieces if p.metadata.get("qa_index", -1) >= 0) == 4


def test_fleet_csv_row_groups_and_summary(estate: Path):
    result = _parse(estate, "fleet-inventory.csv")
    pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename="fleet-inventory.csv")
    assert sum(1 for p in pieces if "row_range" in p.metadata) >= 2  # 60 rows -> multiple groups
    assert any(p.metadata.get("kind") == "table_summary" for p in pieces)


# --------------------------------------------------------------------------- #
# Corpus integrity: the answer key must describe what the documents really contain
# --------------------------------------------------------------------------- #
def _key(estate: Path) -> dict:
    import json

    return json.loads((estate / "answer-key.json").read_text(encoding="utf-8"))


def _evidence_text(estate: Path) -> str:
    texts = []
    for path in sorted(estate.iterdir()):
        if path.name in {"README.md", "ANSWER_KEY.md", "answer-key.json"}:
            continue
        try:
            texts.append("\n".join(p.text for p in parse_file(str(path), path.name).pages))
        except Exception:
            continue  # rejected files (encrypted, legacy .doc) hold no reachable truth
    import zipfile

    with zipfile.ZipFile(estate / "app-portfolio.zip") as zf:
        texts.extend(zf.read(n).decode("utf-8", "replace") for n in zf.namelist())
    return "\n".join(texts).lower()


def test_every_true_server_is_reachable_from_the_documents(estate: Path):
    """Ground truth that no document states would make the answer key unfair."""
    text = _evidence_text(estate)
    key = _key(estate)
    missing = [h for h in key["servers"]["active"] + key["servers"]["aws"] if h not in text]
    assert not missing, f"answer key lists servers no document mentions: {missing}"
    assert all(w in text for w in key["servers"]["fleet_worker_pool"]["hosts"])


def test_every_evidence_file_has_a_documented_trap_or_purpose(estate: Path):
    key = _key(estate)
    trapped = {t["file"] for t in key["traps"]}
    assert trapped <= {p.name for p in estate.iterdir()}, "trap references a file that doesn't exist"
    # Each real-world source file carries at least one planted, documented difficulty.
    for name in (
        "cmdb-export.xlsx",
        "rvtools-export.xlsx",
        "perf-metrics.csv",
        "dependency-connections.csv",
        "integration-register.xlsx",
        "cost-baseline.xlsx",
        "workshop-notes-2026-09-10.md",
        "vendor-handover.txt",
        "dr-plan.docx",
        "app-portfolio.zip",
    ):
        assert name in trapped, name


def test_planted_traps_are_really_in_the_files(estate: Path):
    """Spot-check that the traps the key describes exist as described."""
    from openpyxl import load_workbook

    cmdb = load_workbook(estate / "cmdb-export.xlsx", data_only=True)
    assert cmdb["Lookups"].sheet_state == "hidden"
    total_row = next(r for r in cmdb["Servers"].iter_rows(values_only=True) if r[0] == "TOTAL")
    assert total_row[3] is None  # formula with no cached value reads as empty
    stale = next(r for r in cmdb["Servers"].iter_rows(values_only=True) if r[0] == "ord-api-01")
    assert stale[3] == 8  # CMDB is stale; the truth is 16

    rv = load_workbook(estate / "rvtools-export.xlsx", data_only=True)
    header = [c.value for c in rv["vInfo"][1]]
    row = next(r for r in rv["vInfo"].iter_rows(values_only=True) if r[0] == "ORD-API-01")
    assert row[header.index("CPUs")] == 16 and row[header.index("Memory")] == 32768  # MB, no unit in header
    assert "pay-db-01" not in {r[0] for r in rv["vInfo"].iter_rows(values_only=True)}  # physical: invisible

    perf = (estate / "perf-metrics.csv").read_text(encoding="utf-8").splitlines()
    assert perf[0].count(";") == 6 and "58,4" in "\n".join(perf)

    deps = (estate / "dependency-connections.csv").read_text(encoding="utf-8")
    assert "rpt-01" in deps and ",pay-db-01," in deps and "10.20.5.77" in deps

    assert "ignore all previous instructions" in (estate / "vendor-handover.txt").read_text(encoding="utf-8").lower()


def test_generation_is_deterministic(estate: Path, tmp_path_factory):
    """Re-running the generator must not change the answer key (no randomness, no clock)."""
    spec = importlib.util.spec_from_file_location("_stress_gen_again", _GEN)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    again = tmp_path_factory.mktemp("stress-again")
    module.OUT = again
    module.main()
    assert (again / "answer-key.json").read_text() == (estate / "answer-key.json").read_text()


def test_new_workbooks_parse_without_loss(estate: Path):
    for name in ("rvtools-export.xlsx", "integration-register.xlsx", "cost-baseline.xlsx"):
        result = _parse(estate, name)
        pieces = DocumentChunker().chunk(result.pages, doc_type=result.doc_type, filename=name)
        assert pieces, name
    rv_text = "\n".join(p.text for p in _parse(estate, "rvtools-export.xlsx").pages)
    assert "vInfo" in rv_text and "vHost" in rv_text
