"""End-to-end tests for tables embedded inside a DOCX (architecture/requirements docs) —
verifies document order is preserved, empty cells don't misalign columns, and the
embedded table gets routed through the same shape-guard/summary logic as a whole-sheet
inventory table instead of being flattened into plain prose."""

from docx import Document as DocxDocument

from app.models.entities import DocumentType
from app.services.chunkers import DocumentChunker
from app.services.parsers import TABLE_BLOCK_END, TABLE_BLOCK_START, parse_file


def _build_docx_with_table(path, header, rows, before_text, after_text):
    doc = DocxDocument()
    doc.add_paragraph(before_text)
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    for col, value in enumerate(header):
        table.cell(0, col).text = value
    for r, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            table.cell(r, col).text = value
    doc.add_paragraph(after_text)
    doc.save(str(path))


def test_parse_docx_preserves_document_order_and_marks_table(tmp_path):
    path = tmp_path / "arch.docx"
    _build_docx_with_table(
        path,
        header=["Server", "OS", "vCPU"],
        rows=[["app-01", "RHEL 8", "4"], ["app-02", "Windows Server 2019", "8"]],
        before_text="This section describes the server inventory below.",
        after_text="The above servers are all in production.",
    )
    result = parse_file(str(path), "arch.docx")
    text = result.pages[0].text
    before_idx = text.index("This section describes")
    table_idx = text.index(TABLE_BLOCK_START)
    after_idx = text.index("The above servers are all in production")
    assert before_idx < table_idx < after_idx, "document order was not preserved"
    assert TABLE_BLOCK_END in text


def test_parse_docx_keeps_empty_cells_so_columns_stay_aligned(tmp_path):
    path = tmp_path / "arch.docx"
    _build_docx_with_table(
        path,
        header=["Server", "OS", "vCPU"],
        rows=[["app-01", "RHEL 8", "4"], ["app-02", "", "8"]],  # blank OS cell
        before_text="Intro.",
        after_text="Outro.",
    )
    result = parse_file(str(path), "arch.docx")
    text = result.pages[0].text
    # If the blank cell were dropped instead of kept, "app-02 | 8" would appear (2 fields)
    # instead of "app-02 |  | 8" (3 fields, vCPU still in the 3rd position).
    assert "app-02 |  | 8" in text


def test_embedded_table_is_chunked_with_shape_guard_and_summary(tmp_path):
    path = tmp_path / "arch.docx"
    _build_docx_with_table(
        path,
        header=["Server", "OS", "vCPU"],
        rows=[[f"app-{i}", "RHEL 8", "4"] for i in range(3)],
        before_text="Intro paragraph about the estate.",
        after_text="Closing paragraph about the estate.",
    )
    parsed = parse_file(str(path), "arch.docx")
    pieces = DocumentChunker().chunk(parsed.pages, doc_type=DocumentType.architecture, filename="arch.docx")

    table_pieces = [p for p in pieces if p.metadata.get("embedded_table")]
    summary_pieces = [p for p in table_pieces if p.metadata.get("kind") == "table_summary"]
    row_pieces = [p for p in table_pieces if "row_range" in p.metadata]

    assert table_pieces, "expected the embedded table to be detected and chunked"
    assert len(summary_pieces) == 1
    assert "vcpus: sum=12" in summary_pieces[0].text
    assert row_pieces, "expected row-group chunks for the embedded table"

    # Order preserved: first piece is prose (intro), some middle pieces are the table,
    # last piece is prose (closing) — not all shoved to the end.
    assert "Intro paragraph" in pieces[0].text
    assert "Closing paragraph" in pieces[-1].text
    assert not pieces[0].metadata.get("embedded_table")
    assert not pieces[-1].metadata.get("embedded_table")


def test_embedded_irregular_table_falls_back_and_is_tagged(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("Intro.")
    # A ragged table: build via raw rows of inconsistent length using a normal table
    # then manually break one row's cell count via table-cell text manipulation isn't
    # directly possible with python-docx (it enforces rectangular tables), so we instead
    # exercise this at the chunker level directly, matching how ingest.py would see text
    # that already lost its DOCX table shape (e.g. an unusual manual layout).
    doc.add_paragraph(
        f"{TABLE_BLOCK_START}\n"
        "Server | OS | vCPU\n"
        "app-01 | Linux | 4\n"
        "app-02 | Linux\n"
        "app-03 | Linux | 4 | extra\n"
        f"{TABLE_BLOCK_END}"
    )
    doc.add_paragraph("Outro.")
    path = tmp_path / "arch.docx"
    doc.save(str(path))

    parsed = parse_file(str(path), "arch.docx")
    pieces = DocumentChunker().chunk(parsed.pages, doc_type=DocumentType.architecture, filename="arch.docx")
    assert any(p.metadata.get("shape_guard_failed") for p in pieces)
    assert not any(p.metadata.get("kind") == "table_summary" for p in pieces)
