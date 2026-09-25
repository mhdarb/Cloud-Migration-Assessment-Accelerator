from app.models.entities import Chunk, Document, DocumentType
from app.services.evidence import build_evidence_list, grounded_quote_for_chunk, resolve_evidence


def test_resolve_evidence_returns_human_readable_location(db_session, assessment):
    document = Document(
        assessment_id=assessment.id,
        filename="cmdb-inventory.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        doc_type=DocumentType.inventory,
        storage_path="/tmp/cmdb-inventory.xlsx",
        precedence=100,
    )
    db_session.add(document)
    db_session.flush()
    chunk = Chunk(
        assessment_id=assessment.id,
        document_id=document.id,
        chunk_index=0,
        page=2,
        offset_start=0,
        offset_end=24,
        text="app-01 | 8 | 32 | Linux",
    )
    db_session.add(chunk)
    db_session.commit()

    evidence = resolve_evidence(
        db_session,
        [chunk.id],
        quote="app-01 | 8 | 32 | Linux",
    )

    assert evidence == [
        {
            "chunk_id": chunk.id,
            "document_id": document.id,
            "filename": "cmdb-inventory.xlsx",
            "doc_type": "inventory",
            "page": 2,
            # No sheet/row metadata on this hand-built chunk, and an xlsx "page" is a sheet
            # index, not a page — so no locator rather than a misleading "p. 2".
            "locator": None,
            "quote": "app-01 | 8 | 32 | Linux",
        }
    ]


def test_resolve_evidence_does_not_attach_quote_to_a_chunk_that_lacks_it(db_session, assessment):
    """A claim's `evidence_refs` can list more than one chunk (citations.py sometimes keeps
    every originally-cited chunk when a quote is only grounded in their *combined* text) --
    the quote must not be shown against a chunk that doesn't actually contain it, even
    though the old behavior duplicated it onto every listed chunk unconditionally."""
    document = Document(
        assessment_id=assessment.id,
        filename="architecture.docx",
        content_type="application/docx",
        doc_type=DocumentType.architecture,
        storage_path="/tmp/arch.docx",
        precedence=80,
    )
    db_session.add(document)
    db_session.flush()
    grounded_chunk = Chunk(
        assessment_id=assessment.id,
        document_id=document.id,
        chunk_index=0,
        text="Billing Service is business critical.",
    )
    other_chunk = Chunk(
        assessment_id=assessment.id,
        document_id=document.id,
        chunk_index=1,
        text="Customer Portal has no special criticality noted.",
    )
    db_session.add_all([grounded_chunk, other_chunk])
    db_session.commit()

    evidence = resolve_evidence(
        db_session,
        [grounded_chunk.id, other_chunk.id],
        quote="Billing Service is business critical.",
    )

    by_id = {item["chunk_id"]: item for item in evidence}
    assert by_id[grounded_chunk.id]["quote"] == "Billing Service is business critical."
    assert by_id[other_chunk.id]["quote"] is None


def test_grounded_quote_for_chunk_returns_none_when_not_present():
    entry = {"text": "The server runs RHEL 8."}
    assert grounded_quote_for_chunk(entry, "The server runs Windows Server 2019.") is None
    assert grounded_quote_for_chunk(entry, "The server runs RHEL 8.") == "The server runs RHEL 8."
    assert grounded_quote_for_chunk(entry, None) is None


def test_build_evidence_list_excludes_internal_text_field():
    evidence_by_chunk = {
        "c1": {
            "chunk_id": "c1",
            "document_id": "d1",
            "filename": "f.docx",
            "doc_type": "architecture",
            "page": 1,
            "text": "internal only",
        }
    }
    result = build_evidence_list(evidence_by_chunk, ["c1"], None)
    assert "text" not in result[0]
    assert result[0]["quote"] is None


# --------------------------------------------------------------------------- #
# Format-aware citation locators (instead of a meaningless "p. 1" everywhere)
# --------------------------------------------------------------------------- #
def test_locator_pdf_uses_real_page():
    from app.services.evidence import evidence_locator

    assert evidence_locator("arch.pdf", 4, {}) == "p. 4"


def test_locator_docx_has_no_fake_page():
    """DOCX parses as one unit (no page concept) — never claim "p. 1"."""
    from app.services.evidence import evidence_locator

    assert evidence_locator("architecture-overview.docx", 1, {}) is None


def test_locator_row_range_with_sheet():
    from app.services.evidence import evidence_locator

    meta = {"row_range": [20, 40], "sheet": "Servers"}
    assert evidence_locator("cmdb.xlsx", 1, meta) == 'sheet "Servers", rows 21–40'
    assert evidence_locator("cmdb.xlsx", 1, {"row_range": [2, 3]}) == "row 3"


def test_locator_sheet_summary_question_and_code_file():
    from app.services.evidence import evidence_locator

    assert evidence_locator("cmdb.xlsx", 2, {"sheet": "Legacy Notes"}) == 'sheet "Legacy Notes"'
    assert evidence_locator("cmdb.xlsx", 1, {"kind": "table_summary", "sheet": "Servers"}) == (
        'sheet "Servers", computed summary'
    )
    assert evidence_locator("q.docx", 1, {"qa_index": 2}) == "Q3"
    assert evidence_locator("app.zip", 1, {"file_path": "sample-app/package.json"}) == "sample-app/package.json"


def test_xlsx_sheet_name_reaches_chunk_metadata(tmp_path):
    """End to end through the real parser + chunker: a workbook's sheet name must land on
    each chunk so the citation reads 'sheet "Legacy Notes"', not 'p. 2'."""
    from openpyxl import Workbook

    from app.services.chunkers import DocumentChunker
    from app.services.evidence import evidence_locator
    from app.services.parsers import parse_file

    wb = Workbook()
    ws = wb.active
    ws.title = "Servers"
    ws.append(["hostname", "vcpu"])
    ws.append(["app-01", 4])
    legacy = wb.create_sheet("Legacy Notes")
    legacy.append(["hostname", "os"])
    legacy.append(["app-01", "Windows Server 2016"])
    path = tmp_path / "cmdb-inventory.xlsx"
    wb.save(path)

    parsed = parse_file(str(path), "cmdb-inventory.xlsx")
    pieces = DocumentChunker().chunk(parsed.pages, doc_type=DocumentType.inventory, filename=path.name)
    sheets = {p.metadata.get("sheet") for p in pieces}
    assert {"Servers", "Legacy Notes"} <= sheets
    legacy_piece = next(p for p in pieces if p.metadata.get("sheet") == "Legacy Notes" and "row_range" in p.metadata)
    assert evidence_locator(path.name, legacy_piece.page, legacy_piece.metadata) == 'sheet "Legacy Notes", row 1'
