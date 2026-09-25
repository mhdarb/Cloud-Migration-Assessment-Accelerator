"""Read the questions out of a client questionnaire file, and write answers back into it.

Supported uploads: .xlsx/.xlsm, .csv, .docx, .txt, .md (answered in place — the client
gets their own file back, formatting intact) and .pdf (read-only: answered as an Excel
summary). Every item carries a *locator* recording where the question sits, so the answer
can be written beside it on download.

Security: answers are derived from uploaded documents, i.e. untrusted text. Anything
written into a spreadsheet is forced to a plain string, so a value like "=HYPERLINK(...)"
can never execute as a formula when the client opens the file (CSV/Excel injection).
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.services.questionnaire_extract import normalize_question_key

MAX_QUESTIONS = 500

FORMATS_BY_SUFFIX = {
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".csv": "csv",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
    ".pdf": "pdf",
}
# The download is the client's own format, except PDF (not editable) -> Excel summary.
EXPORT_SUFFIX = {"xlsx": ".xlsx", "csv": ".csv", "docx": ".docx", "txt": ".txt", "md": ".md", "pdf": ".xlsx"}
MEDIA_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".csv": "text/csv; charset=utf-8",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/markdown; charset=utf-8",
}


def questionnaire_format(filename: str) -> str | None:
    return FORMATS_BY_SUFFIX.get(Path(filename).suffix.lower())


@dataclass
class ParsedItem:
    question: str
    locator: dict[str, Any]


@dataclass
class AnswerCell:
    """What gets written back for one question."""

    answer: str
    confidence: float
    needs_review: bool
    sources: str

    @property
    def meta(self) -> str:
        review = "Needs review · " if self.needs_review else ""
        return f"Confidence {round(self.confidence * 100)}% · {review}Sources: {self.sources}"


# --------------------------------------------------------------------------- #
# Question detection
# --------------------------------------------------------------------------- #
# Leading numbering/labels: "Q1:", "Q:", "Question 3.", "1.", "1.2)", "a)", "Q-4 -", and
# section-style refs "A1.", "B2)", "SEC-3:".
_PREFIX = re.compile(
    r"^\s*(?:(?:q(?:uestion)?\s*[-#]?\s*|[a-z]{1,3}-?)?\d+(?:\.\d+)*\s*[:.)\-]"
    r"|q(?:uestion)?\s*[:.)\-]|[a-z][.)])\s+",
    re.I,
)
_INTERROGATIVE = re.compile(
    r"^(what|which|how|does|do|did|is|are|was|were|can|could|will|would|should|shall|who|whom|"
    r"whose|where|when|why|has|have|describe|list|provide|explain|please|identify|outline|"
    r"detail|specify|confirm|state|indicate)\b",
    re.I,
)
_QUESTION_HEADER = re.compile(r"^\s*(questions?|question\s+text|query|queries|requirement)\s*$", re.I)
_ANSWER_HEADER = re.compile(r"^\s*(answers?|responses?|reply|vendor\s+response|client\s+response)\s*$", re.I)
# A blank answer slot in a document template: "Answer:", "Response:", "A:".
_ANSWER_SLOT = re.compile(r"^\s*(answer|response|a)\s*[:.\-]?\s*$", re.I)


def clean_question(text: str) -> str:
    body = _PREFIX.sub("", (text or "").strip(), count=1)
    return " ".join(body.split())


def looks_like_question(text: str, *, in_question_column: bool = False) -> bool:
    raw = (text or "").strip()
    if not raw or len(raw) > 1000:
        return False
    body = clean_question(raw)
    if len(normalize_question_key(body)) < 8:
        return False
    if body.endswith("?"):
        return True
    if in_question_column:
        # The column says these are questions; skip one/two-word section labels.
        return len(body.split()) >= 4
    return bool(_PREFIX.match(raw)) and bool(_INTERROGATIVE.match(body))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _tabular_items(rows: list[list[str]]) -> tuple[int | None, int, int | None, list[tuple[int, str]]]:
    """(header_row, question_col, answer_col, [(row, question)]) for a grid of cells."""
    for r, cells in enumerate(rows[:15]):
        q_cols = [c for c, v in enumerate(cells) if _QUESTION_HEADER.match(v or "")]
        if q_cols:
            q = q_cols[0]
            a = next((c for c, v in enumerate(cells) if c != q and _ANSWER_HEADER.match(v or "")), None)
            items = [
                (i, clean_question(row[q]))
                for i, row in enumerate(rows[r + 1 :], start=r + 1)
                if q < len(row) and looks_like_question(row[q], in_question_column=True)
            ]
            return r, q, a, items
    # No header: the column where most non-empty cells are clearly questions.
    width = max((len(r) for r in rows), default=0)
    best: tuple[int, int] | None = None
    for c in range(width):
        values = [row[c] for row in rows if c < len(row) and row[c]]
        hits = sum(1 for v in values if looks_like_question(v))
        if hits >= 2 and hits * 2 >= len(values) and (best is None or hits > best[1]):
            best = (c, hits)
    if best is None:
        return None, 0, None, []
    q = best[0]
    items = [(i, clean_question(row[q])) for i, row in enumerate(rows) if q < len(row) and looks_like_question(row[q])]
    return None, q, None, items


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def parse_questionnaire(path: str, file_format: str, filename: str = "") -> list[ParsedItem]:
    parser = {
        "xlsx": _parse_xlsx,
        "csv": _parse_csv,
        "docx": _parse_docx,
        "txt": _parse_lines,
        "md": _parse_lines,
        "pdf": _parse_pdf,
    }[file_format]
    items = parser(path, filename)
    # Same question twice in one file: keep the first occurrence.
    seen: set[str] = set()
    unique = []
    for item in items:
        key = normalize_question_key(item.question)
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique[:MAX_QUESTIONS]


def _parse_xlsx(path: str, _filename: str) -> list[ParsedItem]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    out: list[ParsedItem] = []
    for ws in wb.worksheets:
        rows = [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        header, q, a, found = _tabular_items(rows)
        for r, question in found:
            out.append(
                ParsedItem(
                    question,
                    {
                        "kind": "xlsx",
                        "sheet": ws.title,
                        "row": r + 1,
                        "col": q + 1,
                        "answer_col": a + 1 if a is not None else None,
                        "header_row": header + 1 if header is not None else None,
                    },
                )
            )
    wb.close()
    return out


def _read_csv(path: str) -> tuple[list[list[str]], csv.Dialect | type[csv.Dialect]]:
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect: csv.Dialect | type[csv.Dialect] = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)], dialect


def _parse_csv(path: str, _filename: str) -> list[ParsedItem]:
    rows, _ = _read_csv(path)
    header, q, a, found = _tabular_items([[c.strip() for c in r] for r in rows])
    return [
        ParsedItem(question, {"kind": "csv", "row": r, "col": q, "answer_col": a, "header_row": header})
        for r, question in found
    ]


def _parse_docx(path: str, _filename: str) -> list[ParsedItem]:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    out: list[ParsedItem] = []
    paragraphs = doc.paragraphs
    # Walk the body in document order so paragraph and table questions stay interleaved.
    p_index = t_index = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = paragraphs[p_index].text
            if looks_like_question(text):
                nxt = paragraphs[p_index + 1].text if p_index + 1 < len(paragraphs) else ""
                out.append(
                    ParsedItem(
                        clean_question(text),
                        {
                            "kind": "docx_paragraph",
                            "paragraph": p_index,
                            "answer_paragraph": p_index + 1 if _ANSWER_SLOT.match(nxt) else None,
                        },
                    )
                )
            p_index += 1
        elif tag == "tbl":
            table = doc.tables[t_index]
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            _, q, a, found = _tabular_items(rows)
            for r, question in found:
                out.append(
                    ParsedItem(
                        question,
                        {"kind": "docx_table", "table": t_index, "row": r, "col": q, "answer_col": a},
                    )
                )
            t_index += 1
    return out


def _parse_lines(path: str, _filename: str) -> list[ParsedItem]:
    lines = Path(path).read_text(encoding="utf-8-sig", errors="replace").splitlines()
    out: list[ParsedItem] = []
    for i, line in enumerate(lines):
        text = line.lstrip("#>*- \t")  # markdown heading/quote/list markers
        if looks_like_question(text):
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            out.append(
                ParsedItem(
                    clean_question(text),
                    {"kind": "line", "line": i, "answer_line": i + 1 if _ANSWER_SLOT.match(nxt) else None},
                )
            )
    return out


def _parse_pdf(path: str, filename: str) -> list[ParsedItem]:
    from app.services.parsers import parse_file

    parsed = parse_file(path, filename or "questionnaire.pdf")
    out: list[ParsedItem] = []
    for page in parsed.pages:
        for i, line in enumerate((page.text or "").splitlines()):
            if looks_like_question(line):
                out.append(ParsedItem(clean_question(line), {"kind": "pdf", "page": page.page, "line": i}))
    return out


# --------------------------------------------------------------------------- #
# Writing answers back
# --------------------------------------------------------------------------- #
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(value: str) -> str:
    return f"'{value}" if value.startswith(_FORMULA_TRIGGERS) else value


def write_answered(
    file_format: str,
    path: str,
    items: list[tuple[ParsedItem, AnswerCell]],
    *,
    summary: bool = False,
) -> tuple[bytes, str]:
    """(file bytes, suffix). `summary=True` (or a PDF source) returns a fresh Excel table
    of every question and answer instead of filling in the original file."""
    if summary or file_format == "pdf":
        return _write_summary_xlsx(items), ".xlsx"
    writer = {"xlsx": _write_xlsx, "csv": _write_csv, "docx": _write_docx, "txt": _write_lines, "md": _write_lines}
    suffix = Path(path).suffix.lower() if file_format == "xlsx" else EXPORT_SUFFIX[file_format]
    return writer[file_format](path, items, file_format), suffix


def _set_text(cell: Any, value: Any) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # never let a value starting with "=" become a live formula


def _write_xlsx(path: str, items: list[tuple[ParsedItem, AnswerCell]], _fmt: str) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    by_sheet: dict[str, list[tuple[ParsedItem, AnswerCell]]] = defaultdict(list)
    for item, cell in items:
        by_sheet[item.locator["sheet"]].append((item, cell))
    wrap = Alignment(wrap_text=True, vertical="top")
    for sheet, group in by_sheet.items():
        ws = wb[sheet]
        loc = group[0][0].locator
        base = ws.max_column
        answer_col = loc["answer_col"]
        columns: list[tuple[str, int]] = []
        if answer_col is None:
            answer_col = base + 1
            columns.append(("Answer", answer_col))
            base += 1
        columns += [("Confidence", base + 1), ("Needs review", base + 2), ("Sources", base + 3)]
        if loc["header_row"]:
            for label, col in columns:
                header = ws.cell(row=loc["header_row"], column=col)
                _set_text(header, label)
                header.font = Font(bold=True)
        widths = {"Answer": 70, "Sources": 50}
        for label, col in columns:
            if label in widths:
                ws.column_dimensions[get_column_letter(col)].width = widths[label]
        for item, cell in group:
            row = item.locator["row"]
            for value, col in (
                (cell.answer, answer_col),
                (f"{round(cell.confidence * 100)}%", base + 1),
                ("Yes" if cell.needs_review else "No", base + 2),
                (cell.sources, base + 3),
            ):
                target = ws.cell(row=row, column=col)
                _set_text(target, value)
                target.alignment = wrap
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_csv(path: str, items: list[tuple[ParsedItem, AnswerCell]], _fmt: str) -> bytes:
    rows, dialect = _read_csv(path)
    width = max((len(r) for r in rows), default=0)
    loc = items[0][0].locator if items else {"answer_col": None, "header_row": None}
    answer_col = loc["answer_col"]
    extra = ["Confidence", "Needs review", "Sources"]
    if answer_col is None:
        answer_col = width
        extra = ["Answer", *extra]
    new_width = width + len(extra)
    rows = [r + [""] * (new_width - len(r)) for r in rows]
    if loc["header_row"] is not None:
        rows[loc["header_row"]][width:] = extra
    meta_col = new_width - 3  # Confidence, Needs review, Sources are always the last three
    for item, cell in items:
        row = rows[item.locator["row"]]
        row[answer_col] = _csv_safe(cell.answer)
        row[meta_col:] = [
            f"{round(cell.confidence * 100)}%",
            "Yes" if cell.needs_review else "No",
            _csv_safe(cell.sources),
        ]
    out = io.StringIO()
    # Keep only the client's delimiter: a sniffed dialect can come back unable to escape
    # quotes (QUOTE_NONE / no escapechar), and answers routinely contain `"` and `,`.
    csv.writer(out, csv.excel, delimiter=dialect.delimiter).writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def _insert_paragraph_after(paragraph: Any) -> Any:
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _fill_answer(paragraph: Any, cell: AnswerCell) -> Any:
    """Write 'Answer: …' into `paragraph`, then a small metadata line after it."""
    from docx.shared import Pt, RGBColor

    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    label = paragraph.add_run("Answer: ")
    label.bold = True
    paragraph.add_run(cell.answer)
    meta = _insert_paragraph_after(paragraph)
    run = meta.add_run(cell.meta)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    return meta


def _write_docx(path: str, items: list[tuple[ParsedItem, AnswerCell]], _fmt: str) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    paragraphs = doc.paragraphs  # captured before any insertion, so indices stay valid
    tables = doc.tables
    for item, cell in items:
        loc = item.locator
        if loc["kind"] == "docx_paragraph":
            slot = loc.get("answer_paragraph")
            target = paragraphs[slot] if slot is not None else _insert_paragraph_after(paragraphs[loc["paragraph"]])
            _fill_answer(target, cell)
        elif loc["kind"] == "docx_table":
            row = tables[loc["table"]].rows[loc["row"]]
            a = loc.get("answer_col")
            if a is not None and a < len(row.cells):
                target_cell = row.cells[a]
                _fill_answer(target_cell.paragraphs[0], cell)
            else:
                _fill_answer(row.cells[loc["col"]].add_paragraph(), cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_lines(path: str, items: list[tuple[ParsedItem, AnswerCell]], file_format: str) -> bytes:
    lines = Path(path).read_text(encoding="utf-8-sig", errors="replace").splitlines()
    md = file_format == "md"
    by_line = {item.locator["line"]: cell for item, cell in items}
    slots = {item.locator["answer_line"]: cell for item, cell in items if item.locator.get("answer_line") is not None}
    out: list[str] = []
    for i, line in enumerate(lines):
        if i in slots:
            continue  # a blank "Answer:" slot — replaced by the answer block emitted below
        out.append(line)
        if i in by_line:
            cell = by_line[i]
            if md:
                out += ["", f"**Answer:** {cell.answer}", "", f"_{cell.meta}_", ""]
            else:
                out += [f"Answer: {cell.answer}", f"({cell.meta})", ""]
    return ("\n".join(out) + "\n").encode("utf-8")


def _write_summary_xlsx(items: list[tuple[ParsedItem, AnswerCell]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    headers = ["#", "Question", "Answer", "Confidence", "Needs review", "Sources"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for n, (item, cell) in enumerate(items, start=1):
        values = [
            n,
            item.question,
            cell.answer,
            f"{round(cell.confidence * 100)}%",
            "Yes" if cell.needs_review else "No",
            cell.sources,
        ]
        for col, value in enumerate(values, start=1):
            target = ws.cell(row=n + 1, column=col)
            _set_text(target, value)
            target.alignment = wrap
    for letter, width in zip("ABCDEF", (5, 55, 70, 12, 13, 45), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
