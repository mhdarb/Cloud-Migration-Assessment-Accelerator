#!/usr/bin/env python3
"""Generate additional *labeled* eval estates under sample-data/eval/<name>/.

These are original, practical migration scenarios that exercise conditions the Contoso
golden estate does not — heterogeneous headers, header-embedded units (MB), TB-scale
disks, JSON discovery exports keyed under a non-'servers' key, merged datacenter cells,
and hostname-first inventories. Their ground-truth labels live in `eval_datasets.py`; the
eval harness scores extraction/retrieval against them.

Regenerate:  python scripts/generate_eval_estates.py
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document as Docx
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample-data" / "eval"


# --------------------------------------------------------------------------- #
# Estate: Meridian Health (healthcare) — heterogeneous headers, MB-in-header, HIPAA
# --------------------------------------------------------------------------- #
def meridian() -> None:
    base = OUT / "meridian"
    base.mkdir(parents=True, exist_ok=True)

    # Hostname-first CMDB with units in the HEADER (RAM (MB), Disk (GB)) and OS families
    # outside the old 5-family regex (CentOS, Debian).
    (base / "server-inventory.csv").write_text(
        "Hostname,Cores,RAM (MB),Disk (GB),Operating System,Environment\n"
        "mh-web-01,4,16384,200,CentOS 7,prod\n"
        "mh-web-02,4,16384,200,CentOS 7,prod\n"
        "mh-app-01,8,32768,500,RHEL 9,prod\n"
        "mh-db-01,16,65536,2000,Debian 12,prod\n",
        encoding="utf-8",
    )

    doc = Docx()
    doc.add_heading("Meridian Health — Platform Architecture", 0)
    doc.add_paragraph(
        "Application: Patient Portal — patient self-service web application. "
        "Patient Portal depends on Records API."
    )
    doc.add_paragraph(
        "Application: Records API — clinical records service. Records API is hosted on "
        "server mh-app-01 and uses database RecordsDB."
    )
    doc.add_paragraph(
        "Application: Scheduler — appointment scheduling service. Scheduler uses database PortalDB."
    )
    doc.add_paragraph("Server: mh-app-01 runs RHEL 9.")
    doc.add_paragraph("Database: RecordsDB (PostgreSQL) stores clinical records.")
    doc.add_paragraph("Database: PortalDB (SQL Server) stores portal profiles.")
    doc.save(base / "architecture.docx")

    req = Docx()
    req.add_heading("Meridian Wave-1 Requirements & NFR Pack", 0)
    req.add_heading("Compliance", level=1)
    req.add_paragraph("HIPAA controls apply to clinical records stored in RecordsDB.")
    req.add_paragraph("Encryption: TLS 1.2+ in transit and AES-256 at rest.")
    req.add_heading("Non-functional requirements", level=1)
    req.add_paragraph("RTO: 1 hour for Records API.")
    req.add_paragraph("RPO: 5 minutes for RecordsDB.")
    req.add_paragraph("Availability target is 99.95% for production workloads.")
    req.add_paragraph("Latency p95 must stay under 300 ms for record retrieval.")
    req.add_paragraph("Data residency: US only for clinical data.")
    req.add_paragraph("Must support 10,000 concurrent users at peak load.")
    req.save(base / "requirements.docx")


# --------------------------------------------------------------------------- #
# Estate: Atlas Retail — JSON discovery export (TB disks), merged cells, PCI
# --------------------------------------------------------------------------- #
def atlas() -> None:
    base = OUT / "atlas"
    base.mkdir(parents=True, exist_ok=True)

    # Discovery-tool export: records under a non-'servers' key, units in the VALUES,
    # including TB-scale disks.
    payload = {
        "generatedBy": "cloud-discovery-tool",
        "resources": [
            {"name": "atlas-web-01", "vcpu": 4, "memory": "16 GB", "disk": "500 GB", "os": "Ubuntu 22.04"},
            {"name": "atlas-web-02", "vcpu": 4, "memory": "16 GB", "disk": "500 GB", "os": "Ubuntu 22.04"},
            {"name": "atlas-cart-01", "vcpu": 8, "memory": "32 GB", "disk": "1 TB", "os": "RHEL 9"},
            {"name": "atlas-db-01", "vcpu": 16, "memory": "128 GB", "disk": "2 TB", "os": "Windows Server 2022"},
        ],
    }
    (base / "discovery-export.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Servers"
    ws.append(["hostname", "datacenter", "vcpu", "memory_gb", "os"])
    ws.append(["atlas-cache-01", "Region-West", 8, 32, "Ubuntu 22.04"])
    ws.append(["atlas-worker-01", None, 8, 32, "Ubuntu 22.04"])
    ws.merge_cells("B2:B3")  # datacenter merged across the two rows
    wb.save(base / "cmdb.xlsx")

    arch = Docx()
    arch.add_heading("Atlas Retail — Target Architecture", 0)
    arch.add_paragraph(
        "Application: Storefront — customer web storefront. Storefront depends on Cart Service."
    )
    arch.add_paragraph(
        "Application: Cart Service — shopping cart and checkout. Cart Service is hosted on "
        "server atlas-cart-01 and uses database CartDB."
    )
    arch.add_paragraph(
        "Application: Inventory Service — stock and catalog management. Inventory Service "
        "uses database InventoryDB."
    )
    arch.add_paragraph("Database: CartDB (PostgreSQL) stores carts and orders.")
    arch.add_paragraph("Database: InventoryDB (MySQL) stores stock levels.")
    arch.save(base / "architecture.docx")

    q = Docx()
    q.add_heading("Atlas Discovery Questionnaire", 0)
    q.add_paragraph("Q1: Which applications are business critical?")
    q.add_paragraph("Cart Service is business critical during peak sales events.")
    q.add_paragraph("Q2: Are there compliance constraints?")
    q.add_paragraph("PCI-DSS applies to Cart Service payment data.")
    q.add_paragraph("Q3: What is not documented?")
    q.add_paragraph("Firewall rules between regions are not documented in this pack.")
    q.save(base / "questionnaire.docx")


# --------------------------------------------------------------------------- #
# Estate: Orion Financial — unsupported OS (AIX/Solaris), cross-source OS conflict, PCI/SOX
# --------------------------------------------------------------------------- #
def orion() -> None:
    base = OUT / "orion"
    base.mkdir(parents=True, exist_ok=True)

    # Hostname-first CMDB mixing supported and *unsupported* OSes (AIX/Solaris must be
    # blocked, not sized) and a 16-vCPU DB host that exceeds the local catalog.
    (base / "server-inventory.csv").write_text(
        "hostname,vcpu,memory_gb,disk_gb,os,environment\n"
        "orion-web-01,4,16,200,RHEL 9,prod\n"
        "orion-web-02,4,16,200,RHEL 9,prod\n"
        "orion-app-01,8,32,500,Windows Server 2022,prod\n"
        "orion-db-01,16,128,2000,Oracle Linux 8,prod\n"
        "orion-legacy-01,8,32,500,AIX 7.2,prod\n"
        "orion-legacy-02,4,16,200,Solaris 11,prod\n"
        "orion-cache-01,4,16,100,Ubuntu 22.04,prod\n"
        "orion-mq-01,4,16,100,Debian 12,prod\n",
        encoding="utf-8",
    )

    arch = Docx()
    arch.add_heading("Orion Financial — Platform Architecture", 0)
    arch.add_paragraph(
        "Application: Trading Gateway — order routing service. Trading Gateway depends on Risk Engine."
    )
    arch.add_paragraph(
        "Application: Risk Engine — real-time risk scoring. Risk Engine is hosted on server "
        "orion-app-01 and uses database RiskDB."
    )
    arch.add_paragraph(
        "Application: Settlement Service — trade settlement. Settlement Service uses database TradesDB."
    )
    # Deliberate cross-source OS conflict: the CMDB (higher precedence) says Windows Server
    # 2022; this older architecture note says 2019. Reconciliation must pick the CMDB's.
    arch.add_paragraph("Server: orion-app-01 runs Windows Server 2019.")
    arch.add_paragraph("Database: RiskDB (PostgreSQL) stores risk scores.")
    arch.add_paragraph("Database: TradesDB (Oracle) stores settlement records.")
    arch.save(base / "architecture.docx")

    req = Docx()
    req.add_heading("Orion Wave-1 Requirements & NFR Pack", 0)
    req.add_heading("Compliance", level=1)
    req.add_paragraph("PCI-DSS and SOX controls apply to settlement data in TradesDB.")
    req.add_paragraph("Encryption: TLS 1.3 in transit and AES-256 at rest.")
    req.add_heading("Non-functional requirements", level=1)
    req.add_paragraph("RTO: 30 minutes for Trading Gateway.")
    req.add_paragraph("RPO: 5 minutes for TradesDB.")
    req.add_paragraph("Availability target is 99.99% for production workloads.")
    req.add_paragraph("Latency p95 must stay under 150 ms for order routing.")
    req.add_paragraph("Data residency: EU only for settlement data.")
    req.add_paragraph("Must support 20,000 concurrent users at peak load.")
    req.save(base / "requirements.docx")


def main() -> None:
    meridian()
    atlas()
    orion()
    files = sorted(str(p.relative_to(OUT)) for p in OUT.rglob("*") if p.is_file())
    print(f"Wrote {len(files)} eval-estate files under {OUT}:")
    for name in files:
        print(f"  - {name}")


if __name__ == "__main__":
    main()
