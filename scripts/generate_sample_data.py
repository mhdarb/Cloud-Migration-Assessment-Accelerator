#!/usr/bin/env python3
"""Generate synthetic Contoso estate documents."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample-data"


def write_architecture() -> None:
    doc = Document()
    doc.add_heading("Contoso Enterprise Architecture Overview", 0)
    doc.add_paragraph(
        "This document describes the current-state architecture for Contoso's "
        "customer-facing and finance platforms prior to cloud migration assessment."
    )
    doc.add_heading("Applications", level=1)
    doc.add_paragraph(
        "Application: Customer Portal — customer self-service web application. "
        "Customer Portal depends on Billing Service for invoice retrieval and "
        "integrates with Identity Gateway for authentication."
    )
    doc.add_paragraph(
        "Application: Billing Service — core billing and invoice calculation engine. "
        "Billing Service uses Oracle Finance DB and is hosted on server app-bill-01. "
        "Business criticality is high."
    )
    doc.add_paragraph(
        "Application: Identity Gateway — centralized authentication and SSO. "
        "Identity Gateway is hosted on server app-id-01 and uses database AuthDB."
    )
    doc.add_heading("Infrastructure", level=1)
    doc.add_paragraph(
        "Server: app-portal-01 runs Windows Server 2019 and hosts Customer Portal. "
        "Server: app-bill-01 runs RHEL 8 and hosts Billing Service. "
        "Server: app-id-01 runs Ubuntu 22 and hosts Identity Gateway."
    )
    doc.add_paragraph(
        "Database: PortalDB (SQL Server) stores customer profiles. "
        "Database: FinanceDB (Oracle) stores invoices and payment history. "
        "Database: AuthDB (PostgreSQL) stores identity tokens and sessions."
    )
    doc.add_heading("Interfaces", level=1)
    doc.add_paragraph(
        "Interface: Portal-to-Billing REST API on port 8443. "
        "Customer Portal calls Billing Service over HTTPS."
    )
    # Intentional OS conflict with inventory for reconciliation 
    doc.add_paragraph(
        "Note: Server app-bill-01 was previously documented as Windows Server 2016 "
        "in an older runbook; current operations state RHEL 8."
    )
    doc.save(OUT / "architecture-overview.docx")


def write_inventory() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "CMDB Inventory"
    ws.append(
        [
            "Application",
            "Server",
            "Database",
            "Tier",
            "OS",
            "Environment",
            "Owner",
            "vCPU",
            "Memory GB",
            "CPU Utilization %",
            "Memory Utilization %",
            "Disk GB",
            "Disk IOPS",
            "Disk Throughput MBps",
            "Region",
            "Architecture",
        ]
    )
    rows = [
        [
            "Customer Portal",
            "app-portal-01",
            "PortalDB",
            "Web",
            "Windows Server 2019",
            "Prod",
            "Digital",
            4,
            16,
            42,
            58,
            256,
            900,
            100,
            "eastus",
            "x64",
        ],
        [
            "Billing Service",
            "app-bill-01",
            "FinanceDB",
            "App",
            "RHEL 8",
            "Prod",
            "Finance IT",
            8,
            32,
            68,
            72,
            900,
            3200,
            170,
            "eastus",
            "x64",
        ],
        [
            "Identity Gateway",
            "app-id-01",
            "AuthDB",
            "Security",
            "Ubuntu 22",
            "Prod",
            "IAM",
            4,
            16,
            35,
            50,
            128,
            500,
            60,
            "eastus",
            "x64",
        ],
        [
            "Reporting Hub",
            "app-rpt-01",
            "FinanceDB",
            "Analytics",
            "Windows Server 2022",
            "Prod",
            "BI",
            "",
            "",
            "",
            "",
            512,
            1200,
            100,
            "eastus",
            "x64",
        ],
    ]
    for r in rows:
        ws.append(r)
    # Conflicting OS row for app-bill-01 to force reconciliation
    ws2 = wb.create_sheet("Legacy Notes")
    ws2.append(["Server", "OS", "Source"])
    ws2.append(["app-bill-01", "Windows Server 2016", "Legacy CMDB export"])
    wb.save(OUT / "cmdb-inventory.xlsx")


def write_questionnaire() -> None:
    doc = Document()
    doc.add_heading("Migration Assessment Questionnaire", 0)
    doc.add_paragraph("Business & technical Q&A for Contoso wave-1 candidates.")
    qa = [
        (
            "Q: Which applications are business critical?",
            "A: Billing Service and Identity Gateway are business critical. "
            "Customer Portal is medium criticality.",
        ),
        (
            "Q: Are there known compliance constraints?",
            "A: FinanceDB contains PCI-relevant payment metadata and must remain in-region.",
        ),
        (
            "Q: Dependency assumptions",
            "A: Reporting Hub depends on Billing Service batch extracts overnight.",
        ),
        (
            "Q: Known gaps",
            "A: Network topology and firewall rules are not documented in this pack.",
        ),
    ]
    for q, a in qa:
        doc.add_paragraph(q)
        doc.add_paragraph(a)
    doc.save(OUT / "assessment-questionnaire.docx")


def write_requirements_nfr() -> None:
    doc = Document()
    doc.add_heading("Contoso Wave-1 Requirements & NFR Pack", 0)
    doc.add_paragraph("Business requirements document (BRD) for cloud migration readiness.")
    doc.add_heading("Non-functional requirements", level=1)
    doc.add_paragraph("NFR: System must remain available during regional failover drills.")
    doc.add_paragraph("SLA: 99.9% monthly availability for Billing Service customer APIs.")
    doc.add_paragraph("Availability target is 99.9% for production workloads.")
    doc.add_paragraph("RTO: 4 hours for Billing Service.")
    doc.add_paragraph("RPO: 15 minutes for FinanceDB.")
    doc.add_paragraph("Latency p95 must stay under 250 ms for invoice retrieval.")
    doc.add_paragraph("Data residency: EU and US in-region only for payment metadata.")
    doc.add_paragraph("Must support 5,000 concurrent users at peak load during month-end close.")
    doc.add_heading("Compliance", level=1)
    doc.add_paragraph("PCI-DSS controls apply to payment metadata stored in FinanceDB.")
    doc.add_paragraph("Encryption: TLS 1.2+ in transit and AES-256 at rest for databases.")
    doc.add_heading("Integration contracts", level=1)
    doc.add_paragraph("Billing Service integrates with Identity Gateway and emits events to Kafka.")
    doc.add_paragraph("Constraints: mainframe batch window cannot move before Q2.")
    doc.save(OUT / "requirements-nfr.docx")


def write_sample_app_zip() -> None:
    import zipfile

    app_dir = OUT / "sample-app"
    app_dir.mkdir(parents=True, exist_ok=True)
    (app_dir / "package.json").write_text(
        """{
  "name": "contoso-billing-api",
  "version": "1.2.0",
  "engines": { "node": ">=18" },
  "dependencies": {
    "express": "^4.19.0",
    "pg": "^8.12.0",
    "ioredis": "^5.4.1",
    "kafkajs": "^2.2.4"
  }
}
""",
        encoding="utf-8",
    )
    (app_dir / "Dockerfile").write_text(
        "FROM node:18-alpine\nWORKDIR /app\nCOPY package.json ./\n"
        "RUN npm install --omit=dev\nCOPY . .\nEXPOSE 8080\nCMD [\"node\", \"server.js\"]\n",
        encoding="utf-8",
    )
    (app_dir / "docker-compose.yml").write_text(
        "services:\n  api:\n    build: .\n    ports:\n      - \"8080:8080\"\n"
        "  postgres:\n    image: postgres:16\n  redis:\n    image: redis:7\n",
        encoding="utf-8",
    )
    (app_dir / ".env.example").write_text(
        "DATABASE_URL=\nREDIS_URL=\nKAFKA_BROKERS=\nPORT=8080\n", encoding="utf-8"
    )
    (app_dir / "server.js").write_text(
        'console.log("contoso-billing-api");\n', encoding="utf-8"
    )
    zip_path = OUT / "sample-app.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in app_dir.rglob("*"):
            if path.is_file():
                zf.write(path, arcname=str(path.relative_to(OUT)))


def write_readme() -> None:
    (OUT / "README.md").write_text(
        """# Sample data — Contoso estate

Synthetic documents:

| File | Type | Purpose |
|------|------|---------|
| `architecture-overview.docx` | Architecture | Apps, servers, DBs, interfaces |
| `cmdb-inventory.xlsx` | Inventory | CMDB export with sizing metrics, one missing-data case, and an OS conflict |
| `assessment-questionnaire.docx` | Questionnaire | Criticality, compliance, gaps |
| `requirements-nfr.docx` | Requirements / NFR | SLA, RTO/RPO, PCI, residency |
| `sample-app.zip` | Code snapshot | Node billing API manifests (pg/redis/kafka) |

Regenerate with:

```bash
python scripts/generate_sample_data.py
```
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_architecture()
    write_inventory()
    write_questionnaire()
    write_requirements_nfr()
    write_sample_app_zip()
    write_readme()
    print(f"Wrote sample documents to {OUT}")


if __name__ == "__main__":
    main()
