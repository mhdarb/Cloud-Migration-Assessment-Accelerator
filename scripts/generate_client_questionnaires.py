"""Generate client questionnaires for trying the Questions-tab upload/download feature.

They are written the way clients actually send them — an Excel RFI with title rows,
section dividers and a "Vendor Response" column; a Word security review mixing numbered
questions, blank "Answer:" slots and a table; a plain CSV — and ask about the Zephyr
Logistics estate in `sample-data/stress/`, so answers can be grounded when that estate is
the assessment's evidence.

    python scripts/generate_client_questionnaires.py
"""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parents[1] / "sample-data" / "questionnaires"

RFI: list[tuple[str, list[str]]] = [
    (
        "Estate & inventory",
        [
            "Which applications, servers and databases are in scope for wave 1?",
            "Which operating systems and versions run on the in-scope servers?",
            "How many vCPUs and how much memory does the order database host have?",
        ],
    ),
    (
        "Resilience",
        [
            "What are the RPO and RTO targets for the Order API?",
            "Is failover between DC-North and DC-South automated or manual?",
            "What availability SLA applies to the Payments Service?",
        ],
    ),
    (
        "Security & compliance",
        [
            "Which workloads are in PCI DSS scope?",
            "Are there data residency constraints on EU order data?",
            "Is data encrypted at rest and in transit?",
        ],
    ),
    (
        "Dependencies & blockers",
        [
            "Which services does the Order API depend on?",
            "What technical debt or blockers could delay migration?",
            "Which integrations use message queues or Kafka?",
        ],
    ),
]


def _rfi_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RFI"
    ws["A1"] = "Zephyr Logistics — Cloud Migration RFI"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Please complete the Vendor Response column. Return by 15 Oct."
    ws.append([])
    header = ["Ref", "Domain", "Question", "Vendor Response", "Comments"]
    ws.append(header)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    ref = 0
    for domain, questions in RFI:
        ws.append(["", domain.upper()])  # section divider: not a question
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        for question in questions:
            ref += 1
            ws.append([f"Z-{ref:02d}", domain, question, None, None])
    for letter, width in zip("ABCDE", (8, 22, 70, 60, 30), strict=True):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    notes = wb.create_sheet("Instructions")
    notes["A1"] = "Answer every question. Mark anything you cannot evidence as 'TBC'."
    wb.save(path)


def _security_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Zephyr Logistics — Security Review Questionnaire", 1)
    doc.add_paragraph("Prepared by the Zephyr CISO office for migration partners.")
    doc.add_heading("Part A — Data protection", 2)
    for n, question in enumerate(
        [
            "Which workloads process cardholder data?",
            "Where must EU customer data be stored?",
            "Describe how secrets and credentials are managed today",
        ],
        start=1,
    ):
        doc.add_paragraph(f"A{n}. {question}")
        doc.add_paragraph("Answer:")
    doc.add_heading("Part B — Operations", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("#", "Question", "Response"), strict=True):
        cell.text = label
    for n, question in enumerate(
        [
            "How are backups taken and how often are restores tested?",
            "Is failover between data centres manual?",
            "Which components are single points of failure?",
        ],
        start=1,
    ):
        row = table.add_row().cells
        row[0].text, row[1].text = f"B{n}", question
    doc.add_paragraph("Thank you for completing this questionnaire.")
    doc.save(path)


def _quick_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "question", "answer"])
        writer.writerow([1, "What runtime does fleet-ingestor use?", ""])
        writer.writerow([2, "Which database does the Payments service use?", ""])
        writer.writerow([3, "Is the Order API containerised?", ""])


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    _rfi_xlsx(OUT / "zephyr-client-rfi.xlsx")
    _security_docx(OUT / "zephyr-security-review.docx")
    _quick_csv(OUT / "zephyr-quick-questions.csv")
    for p in sorted(OUT.iterdir()):
        print(p.relative_to(OUT.parents[1]))


if __name__ == "__main__":
    main()
